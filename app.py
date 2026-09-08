import io
import os
import streamlit as st
import pandas as pd
import plotly.express as px
from dotenv import load_dotenv
from sqlalchemy import create_engine
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

st.set_page_config(
    page_title="Crypto Market Telemetry & ETL Monitor",
    page_icon=":material/monetization_on:",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Custom Design System: #241C40 (Base), #CDF27E (Accent), #6F4BF2 (Interactive), #A38DF2 (Text/Muted)
st.markdown("""
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Material+Symbols+Outlined:opsz,wght,FILL,GRAD@20..48,100..700,0..1,-50..200" />
<style>
    /* Hide Default Streamlit Chrome (Top Bar, Deploy Button, Menu, Footer) */
    #MainMenu {visibility: hidden !important; display: none !important;}
    header[data-testid="stHeader"] {visibility: hidden !important; display: none !important;}
    [data-testid="stToolbar"] {visibility: hidden !important; display: none !important;}
    .stDeployButton {display: none !important;}
    footer {visibility: hidden !important;}

    /* Base App Canvas */
    .stApp {
        background-color: #241C40 !important;
        color: #FFFFFF !important;
    }

    /* Top Bar Header Typography */
    .nav-title {
        color: #CDF27E !important;
        font-size: 2.2rem;
        font-weight: 800;
        margin: 0 !important;
        line-height: 1.15 !important;
        letter-spacing: -0.02em;
        display: flex;
        align-items: center;
        gap: 10px;
    }
    .nav-subtitle {
        color: #A38DF2 !important;
        font-size: 0.95rem;
        margin: 6px 0 0 0 !important;
        letter-spacing: 0.02em;
    }

    /* Top Action Buttons (Sync) */
    .stButton > button {
        background-color: #CDF27E !important;
        color: #241C40 !important;
        font-weight: 800 !important;
        font-size: 0.88rem !important;
        border: none !important;
        border-radius: 20px !important;
        padding: 8px 18px !important;
        box-shadow: 0 4px 14px rgba(205, 242, 126, 0.25) !important;
        transition: all 0.2s ease-in-out;
    }
    .stButton > button:hover {
        background-color: #6F4BF2 !important;
        color: #CDF27E !important;
        transform: translateY(-1px);
        box-shadow: 0 6px 18px rgba(111, 75, 242, 0.4) !important;
    }

    /* Top Action Buttons (Export XLSX) */
    .stDownloadButton > button {
        background-color: #1a1430 !important;
        color: #CDF27E !important;
        font-weight: 800 !important;
        font-size: 0.88rem !important;
        border: 1.5px solid #6F4BF2 !important;
        border-radius: 20px !important;
        padding: 8px 18px !important;
        box-shadow: 0 4px 14px rgba(0, 0, 0, 0.25) !important;
        transition: all 0.2s ease-in-out;
    }
    .stDownloadButton > button:hover {
        background-color: #6F4BF2 !important;
        color: #FFFFFF !important;
        border-color: #CDF27E !important;
        transform: translateY(-1px);
        box-shadow: 0 6px 18px rgba(111, 75, 242, 0.4) !important;
    }

    /* Divider */
    .nav-divider {
        border-bottom: 1px solid rgba(163, 141, 242, 0.2);
        margin: 14px 0 24px 0;
    }

    /* KPI Metric Cards */
    [data-testid="stMetric"] {
        background-color: #1a1430 !important;
        border: 1px solid #6F4BF2 !important;
        border-top: 4px solid #CDF27E !important;
        border-radius: 12px;
        padding: 14px 18px;
        box-shadow: 0 4px 14px rgba(0, 0, 0, 0.25);
    }
    [data-testid="stMetricLabel"] {
        color: #A38DF2 !important;
        font-weight: 700;
        font-size: 0.78rem;
        letter-spacing: 0.08em;
        text-transform: uppercase;
    }
    [data-testid="stMetricValue"] {
        color: #CDF27E !important;
        font-weight: 800;
        font-size: 1.75rem;
    }
    [data-testid="stMetricDelta"] div {
        color: #CDF27E !important;
        font-weight: 700;
    }

    /* Side Cards */
    .side-card {
        background-color: #1a1430;
        border: 1px solid #6F4BF2;
        border-left: 5px solid #CDF27E;
        border-radius: 12px;
        padding: 16px 20px;
        box-shadow: 0 4px 14px rgba(0, 0, 0, 0.25);
        margin-bottom: 14px;
    }
    .side-card-title {
        color: #CDF27E;
        font-weight: 800;
        font-size: 0.85rem;
        letter-spacing: 0.06em;
        text-transform: uppercase;
        margin: 0 0 10px 0;
    }
    .side-card p {
        color: #FFFFFF;
        margin: 0 0 8px 0;
        font-size: 0.88rem;
    }
    .side-card b {
        color: #A38DF2;
    }

    /* Section Headings */
    h3 {
        color: #CDF27E !important;
        font-weight: 800 !important;
        font-size: 1.2rem !important;
        margin: 0 !important;
    }

    /* Table Badge */
    .table-badge {
        background-color: rgba(205, 242, 126, 0.12);
        color: #CDF27E !important;
        border: 1.5px solid #CDF27E;
        border-radius: 6px;
        padding: 2px 10px;
        font-family: monospace;
        font-size: 0.9rem;
        font-weight: 700;
        box-shadow: 0 0 10px rgba(205, 242, 126, 0.15);
    }

    /* Dataframe Box */
    [data-testid="stDataFrame"] {
        border: 2px solid #6F4BF2 !important;
        border-top: 4px solid #CDF27E !important;
        border-radius: 12px !important;
        background-color: #1a1430 !important;
        box-shadow: 0 6px 20px rgba(0, 0, 0, 0.35);
    }

    /* Center Table Headers */
    [data-testid="stDataFrame"] div[role="columnheader"] {
        color: #CDF27E !important;
        font-weight: 800 !important;
        letter-spacing: 0.04em;
        justify-content: center !important;
        text-align: center !important;
    }
    [data-testid="stDataFrame"] div[role="columnheader"] > div {
        justify-content: center !important;
        text-align: center !important;
        width: 100% !important;
    }

    /* Center Table Cells */
    [data-testid="stDataFrame"] div[role="gridcell"] {
        justify-content: center !important;
        text-align: center !important;
    }

    /* Floating Toolbar Styling */
    [data-testid="stElementToolbar"] {
        opacity: 1 !important;
        visibility: visible !important;
        display: flex !important;
        top: -34px !important;
        right: 0px !important;
        background: transparent !important;
        border: none !important;
        box-shadow: none !important;
        padding: 0 !important;
        z-index: 999999 !important;
    }
    [data-testid="stElementToolbar"] button {
        color: #CDF27E !important;
        background: transparent !important;
        border: none !important;
    }
    [data-testid="stElementToolbar"] button:hover {
        color: #241C40 !important;
        background: #CDF27E !important;
        border-radius: 6px !important;
    }

    /* Hide the built-in raw CSV download button */
    [data-testid="stElementToolbar"] button[aria-label="Download as CSV"],
    [data-testid="stElementToolbar"] button[title="Download as CSV"] {
        display: none !important;
    }
</style>
""", unsafe_allow_html=True)

# Helper function to read from st.secrets (Cloud) or os.getenv (Local)
def get_config(key: str, default: str = "") -> str:
    try:
        if hasattr(st, "secrets") and key in st.secrets:
            return str(st.secrets[key])
    except Exception:
        pass
    return os.getenv(key, default)

# Database Engine
@st.cache_resource
def get_db_engine():
    user = get_config("DB_USER", "postgres")
    password = get_config("DB_PASSWORD", "etl_password")
    host = get_config("DB_HOST", "localhost")
    port = get_config("DB_PORT", "5432")
    database = get_config("DB_NAME", "postgres")
    
    # Use sslmode=require for Supabase pooler connections
    return create_engine(
        f"postgresql://{user}:{password}@{host}:{port}/{database}?sslmode=require"
    )

engine = get_db_engine()

# Styled Excel Exporter
def generate_styled_excel(df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    export_df = df.copy()

    # Strip timezone to prevent openpyxl ValueError
    if "last_snapshot_time" in export_df.columns:
        export_df["last_snapshot_time"] = pd.to_datetime(export_df["last_snapshot_time"]).dt.tz_localize(None)

    export_df = export_df.rename(columns={
        "coin_id": "Coin Name",
        "symbol": "Ticker",
        "latest_price": "Latest Price",
        "rolling_avg_price": "Rolling Avg Price",
        "rolling_volatility": "Volatility (σ)",
        "volume_to_mcap_ratio": "Vol / MCap",
        "last_snapshot_time": "Last Snapshot (UTC)"
    })

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export_df.to_excel(writer, sheet_name="Market Telemetry", index=False)
        workbook = writer.book
        worksheet = writer.sheets["Market Telemetry"]

        # Design system palette
        c_header_bg = "241C40"
        c_header_font = "CDF27E"
        c_data_bg = "A38DF2"
        c_data_font = "241C40"
        c_border = "6F4BF2"

        # Styles
        header_fill = PatternFill(start_color=c_header_bg, end_color=c_header_bg, fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color=c_header_font)

        data_fill = PatternFill(start_color=c_data_bg, end_color=c_data_bg, fill_type="solid")
        data_font = Font(name="Segoe UI", size=10, bold=True, color=c_data_font)

        thin_border = Border(
            left=Side(style="thin", color=c_border),
            right=Side(style="thin", color=c_border),
            top=Side(style="thin", color=c_border),
            bottom=Side(style="thin", color=c_border)
        )
        center_align = Alignment(horizontal="center", vertical="center")

        # Format Header Row (First Line)
        worksheet.row_dimensions[1].height = 28
        for col_idx in range(1, len(export_df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        # Format Data Rows
        for row_idx in range(2, len(export_df) + 2):
            worksheet.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(export_df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.fill = data_fill
                cell.font = data_font
                cell.alignment = center_align
                cell.border = thin_border

                val = cell.value
                if isinstance(val, (int, float)):
                    if col_idx in [3, 4]:
                        cell.number_format = '"$"#,##0.0000'
                    elif col_idx == 5:
                        cell.number_format = '0.0000'
                    elif col_idx == 6:
                        cell.number_format = '0.000000'

        # Auto-fit Column Widths
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 16)

    return output.getvalue()

# Queries
@st.cache_data(ttl=30)
def load_view_metrics():
    query = """
    SELECT 
        v.coin_id,
        v.symbol,
        v.latest_price,
        v.rolling_avg_price,
        v.rolling_volatility,
        v.volume_to_mcap_ratio,
        v.last_snapshot_time
    FROM view_crypto_rolling_metrics v
    JOIN crypto_price_history h 
      ON v.symbol = h.symbol 
     AND v.last_snapshot_time = h.recorded_at
    ORDER BY h.market_cap DESC;
    """
    return pd.read_sql(query, engine)

@st.cache_data(ttl=30)
def load_historical_series(symbol: str):
    query = f"""
    SELECT 
        symbol,
        price,
        market_cap,
        total_volume,
        recorded_at
    FROM crypto_price_history
    WHERE symbol = '{symbol}'
    ORDER BY recorded_at ASC;
    """
    return pd.read_sql(query, engine)

view_df = load_view_metrics()

# Top Navbar with Crypto / Currency Icon
nav_col_left, nav_col_sync, nav_col_export = st.columns([3.0, 0.9, 0.9], vertical_alignment="center")

with nav_col_left:
    st.markdown("""
        <div class="nav-title">
            <span class="material-symbols-outlined" style="font-size: 2.2rem; color: #CDF27E; line-height: 1;">monetization_on</span>
            Crypto Pipeline Telemetry
        </div>
        <p class="nav-subtitle">Automated Ingestion • PostgreSQL Window Analytics • Immutable Time-Series</p>
    """, unsafe_allow_html=True)

with nav_col_sync:
    if st.button("Sync Telemetry", icon=":material/sync:", width="stretch"):
        st.cache_data.clear()

with nav_col_export:
    excel_data = generate_styled_excel(view_df)
    st.download_button(
        label="Export Report",
        icon=":material/download:",
        data=excel_data,
        file_name="crypto_telemetry_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch"
    )

st.markdown('<div class="nav-divider"></div>', unsafe_allow_html=True)

if not view_df.empty:
    # Dynamic Top 3 from view_df ranking
    cols = st.columns(4)

    for i in range(min(3, len(view_df))):
        row = view_df.iloc[i]
        diff = row['latest_price'] - row['rolling_avg_price']
        coin_label = f"{row['coin_id'].replace('-', ' ').title()} ({row['symbol']})"
        
        with cols[i]:
            st.metric(
                label=coin_label,
                value=f"${row['latest_price']:,.2f}" if row['latest_price'] >= 1 else f"${row['latest_price']:,.4f}",
                delta=f"${diff:+,.2f} vs avg" if abs(diff) >= 0.01 else f"${diff:+,.4f} vs avg"
            )

    with cols[3]:
        st.metric(
            label="Pipeline Status",
            value="ACTIVE 0x0",
            delta=f"{len(view_df)} Assets Tracking"
        )

    st.write("")

    # Main Row: Chart on Left vs Stacked Cards on Right
    chart_col, side_col = st.columns([2.4, 1.0])

    with chart_col:
        symbol_name_map = dict(
            zip(
                view_df["symbol"],
                view_df["coin_id"].str.replace("-", " ").str.title() + " (" + view_df["symbol"] + ")"
            )
        )

        head_col, select_col = st.columns([2.0, 1.2])
        with head_col:
            st.markdown("### Price Action Across Pipeline Executions")
        with select_col:
            selected_symbol = st.selectbox(
                "Select Asset:",
                options=view_df["symbol"].tolist(),
                format_func=lambda sym: symbol_name_map.get(sym, sym),
                index=0,
                label_visibility="collapsed"
            )

        hist_df = load_historical_series(selected_symbol)
        if len(hist_df) > 1:
            fig = px.area(
                hist_df,
                x="recorded_at",
                y="price",
                labels={"price": "USD Price", "recorded_at": "Execution Run"},
            )
            fig.update_traces(
                line_color="#CDF27E",
                fillcolor="rgba(111, 75, 242, 0.40)",
                line_width=3,
                mode="lines+markers",
                marker=dict(size=8, color="#CDF27E", line=dict(width=2, color="#241C40"))
            )
            y_min = hist_df["price"].min() * 0.998
            y_max = hist_df["price"].max() * 1.002
            fig.update_layout(
                margin=dict(l=15, r=15, t=10, b=15),
                height=385,
                xaxis=dict(showgrid=False, color="#A38DF2"),
                yaxis=dict(
                    showgrid=True, 
                    gridcolor="rgba(163, 141, 242, 0.15)", 
                    color="#A38DF2",
                    range=[y_min, y_max],
                    tickformat="$,.2f"
                ),
                plot_bgcolor="#1a1430",
                paper_bgcolor="#1a1430",
                font=dict(color="#FFFFFF")
            )
            st.plotly_chart(fig, width="stretch")
        else:
            st.info(f"Only 1 snapshot logged for {selected_symbol}. Consecutive runs will plot here.")

    with side_col:
        selected_stats = view_df[view_df["symbol"] == selected_symbol].iloc[0]
        formatted_coin_label = symbol_name_map.get(selected_symbol, selected_symbol)
        
        # Asset Telemetry Card
        st.markdown(f"""
        <div class="side-card">
            <div class="side-card-title">{formatted_coin_label} Telemetry</div>
            <p><b>Spot Price:</b> <span style="font-weight: 800; font-size: 1.05rem; color: #CDF27E;">${selected_stats['latest_price']:,.4f}</span></p>
            <p><b>Rolling 24-Run Avg:</b> ${selected_stats['rolling_avg_price']:,.4f}</p>
            <p><b>Volatility (STDDEV):</b> <span style="color: #CDF27E; font-weight: 700;">{selected_stats['rolling_volatility']:,.4f}</span></p>
            <p><b>Vol / MCap Ratio:</b> {selected_stats['volume_to_mcap_ratio']:,.6f}</p>
            <p style="margin-bottom: 0;"><b>Last Sync:</b> {selected_stats['last_snapshot_time'].strftime('%H:%M:%S UTC')}</p>
        </div>
        """, unsafe_allow_html=True)

        # Pipeline Health Card
        st.markdown("""
        <div class="side-card" style="border-left-color: #6F4BF2;">
            <div class="side-card-title" style="color: #A38DF2;">Pipeline Engine Health</div>
            <p><b>Orchestrator:</b> GitHub Actions (Cron */10)</p>
            <p><b>Target Database:</b> Supabase PostgreSQL (Cloud)</p>
            <p><b>Circuit Breaker:</b> Active (&lt; 50% Δ assertion)</p>
            <p><b>Ingestion Status:</b> Automated 24/7</p>
            <p style="margin-bottom: 0;"><b>Idempotency:</b> ON CONFLICT DO NOTHING</p>
        </div>
        """, unsafe_allow_html=True)

    st.write("")

    # System Analytical View Table Header
    st.markdown(
        '<div style="margin-bottom: 12px;">'
        '### System Analytical View <span class="table-badge">view_crypto_rolling_metrics</span>'
        '</div>', 
        unsafe_allow_html=True
    )

    st.dataframe(
        view_df,
        column_config={
            "coin_id": st.column_config.TextColumn("Coin Name"),
            "symbol": st.column_config.TextColumn("Ticker"),
            "latest_price": st.column_config.NumberColumn("Latest Price", format="$%.4f"),
            "rolling_avg_price": st.column_config.NumberColumn("Rolling Avg Price", format="$%.4f"),
            "rolling_volatility": st.column_config.NumberColumn("Volatility (σ)", format="%.4f"),
            "volume_to_mcap_ratio": st.column_config.NumberColumn("Vol / MCap", format="%.6f"),
            "last_snapshot_time": st.column_config.DatetimeColumn("Last Snapshot (UTC)", format="YYYY-MM-DD HH:mm:ss")
        },
        width="stretch",
        hide_index=True
    )